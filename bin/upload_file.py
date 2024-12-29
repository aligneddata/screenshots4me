import requests
import sys
import os
from PyPDF2 import PdfReader
from docx import Document
import pytesseract
from PIL import Image
import openpyxl

def convert_to_text(filepath):
    if filepath.endswith('.pdf'):
        reader = PdfReader(filepath)
        return ''.join([page.extract_text() for page in reader.pages])
    elif filepath.endswith('.docx'):
        doc = Document(filepath)
        return '\n'.join([para.text for para in doc.paragraphs])
    elif filepath.endswith('.xlsx'):
        wb = openpyxl.load_workbook(filepath)
        sheet = wb.active
        return '\n'.join([' '.join([str(cell.value) for cell in row]) for row in sheet.rows])
    elif filepath.endswith(('.png', '.jpg', '.jpeg')):
        return pytesseract.image_to_string(Image.open(filepath))
    else:
        raise ValueError("Unsupported file type")

def upload_to_api(filename, filetype, text, filepath):
    url = "http://localhost:8080/api/documents/upload"
    payload = {
        "filename": filename,
        "fileType": filetype,
        "convertedText": text,
        "filePath": filepath
    }
    response = requests.post(url, data=payload)
    print(response.text)

if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python script.py <filepath>")
        sys.exit(1)
    filepath = sys.argv[1]
    filename = os.path.basename(filepath)
    filetype = filename.split('.')[-1].upper()
    text = convert_to_text(filepath)
    upload_to_api(filename, filetype, text, filepath)